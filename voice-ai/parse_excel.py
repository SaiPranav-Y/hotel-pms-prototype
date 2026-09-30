"""Parse the Karivena Excel file into structured knowledge base data."""
import openpyxl
import json
from pathlib import Path

EXCEL_PATH = Path(__file__).parent / "Karivenam_Consolidated Present Room Details_28.04.2026_ Changes.xlsx"
OUTPUT_PATH = Path(__file__).parent / "app" / "room_data.json"


def parse():
    wb = openpyxl.load_workbook(str(EXCEL_PATH))
    all_locations = {}

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        location_key = sheet_name.lower().replace(" ", "_")

        # Find header row
        header_row = None
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), 1):
            if row and any(str(c).strip().lower() == "room no" for c in row if c):
                header_row = i
                break

        if not header_row:
            all_locations[location_key] = {
                "name": sheet_name,
                "total_rooms": 0,
                "ac_rooms": 0,
                "nonac_rooms": 0,
                "buildings": [],
                "prices_ac": [],
                "prices_nonac": [],
                "rooms": [],
            }
            continue

        rooms = []
        ac_count = 0
        nonac_count = 0
        prices_ac = set()
        prices_nonac = set()
        buildings = set()

        for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
            if not row[0] or not str(row[0]).strip().replace(".", "").isdigit():
                continue

            branch = str(row[1]).strip() if row[1] else sheet_name
            building = str(row[2]).strip() if row[2] else ""
            floor = str(row[3]).strip() if row[3] else ""
            room_no = str(row[4]).strip() if row[4] else ""
            room_type = str(row[5]).strip() if row[5] else ""
            occupancy = str(row[6]).strip() if row[6] else ""
            price = row[7] if row[7] else 0

            try:
                price = int(float(price))
            except (ValueError, TypeError):
                price = 0

            is_ac = "ac" in room_type.lower() and "non" not in room_type.lower()
            is_nonac = "non" in room_type.lower()

            if is_ac:
                ac_count += 1
                if price > 0:
                    prices_ac.add(price)
            elif is_nonac:
                nonac_count += 1
                if price > 0:
                    prices_nonac.add(price)

            if building:
                buildings.add(building)

            rooms.append({
                "room_no": room_no,
                "building": building,
                "floor": floor,
                "type": "AC" if is_ac else "Non-AC",
                "occupancy": occupancy,
                "price": price,
            })

        all_locations[location_key] = {
            "name": sheet_name,
            "total_rooms": len(rooms),
            "ac_rooms": ac_count,
            "nonac_rooms": nonac_count,
            "buildings": sorted(buildings),
            "prices_ac": sorted(prices_ac),
            "prices_nonac": sorted(prices_nonac),
            "rooms": rooms,
        }

    return all_locations


def save_summary(all_locations):
    """Save summary (without individual rooms) for quick loading."""
    summary = {}
    for key, loc in all_locations.items():
        summary[key] = {
            "name": loc["name"],
            "total_rooms": loc["total_rooms"],
            "ac_rooms": loc["ac_rooms"],
            "nonac_rooms": loc["nonac_rooms"],
            "buildings": loc["buildings"],
            "prices_ac": loc["prices_ac"],
            "prices_nonac": loc["prices_nonac"],
        }
    with open(str(OUTPUT_PATH), "w") as f:
        json.dump(summary, f, indent=2)
    return summary


if __name__ == "__main__":
    print("Parsing Karivena Excel file...")
    data = parse()

    print("\nLOCATION SUMMARY:")
    print("=" * 75)
    total_rooms = 0
    for key, loc in data.items():
        total_rooms += loc["total_rooms"]
        ac_p = loc["prices_ac"] if loc["prices_ac"] else ["free"]
        nonac_p = loc["prices_nonac"] if loc["prices_nonac"] else ["free"]
        print(
            f"  {loc['name']:15} | "
            f"Total: {loc['total_rooms']:3} | "
            f"AC: {loc['ac_rooms']:3} (INR {ac_p}) | "
            f"Non-AC: {loc['nonac_rooms']:3} (INR {nonac_p})"
        )

    print(f"\n  TOTAL ROOMS: {total_rooms}")
    print(f"  LOCATIONS: {len(data)}")

    summary = save_summary(data)
    print(f"\n  Saved to: {OUTPUT_PATH}")
